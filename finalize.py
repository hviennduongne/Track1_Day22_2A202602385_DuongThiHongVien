from pathlib import Path
import json, math, zipfile, copy, xml.etree.ElementTree as ET
import openpyxl
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
import pypdfium2

P=Path(__file__).parent; O=P/'output';Q=P/'qa';O.mkdir(exist_ok=True)
inputs=json.loads((P/'inputs.json').read_text(encoding='utf8'));layout=json.loads((P/'layout.json').read_text())
source=P/'templates/model.xlsx';export=Q/'artifact-export.xlsx';dest=O/'DuongThiHongVien_Day22_model.xlsx'
src=openpyxl.load_workbook(source);ex=openpyxl.load_workbook(export,data_only=True)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};N='{'+ns['m']+'}'
ET.register_namespace('',ns['m']);ET.register_namespace('r','http://schemas.openxmlformats.org/officeDocument/2006/relationships')
repair=[]
raw=openpyxl.load_workbook(export)
for a,b in zip(src,raw):
 if len(a.data_validations.dataValidation)!=len(b.data_validations.dataValidation):repair.append(a.title+': export validation count differs')
 if any(c.value!=b[c.coordinate].value for row in a for c in row if c.data_type=='f'):repair.append(a.title+': export formula text differs')
# Preserve original OOXML package/features. Only artifact-authored input values and
# calculated formula caches are transplanted; original formula text is retained.
with zipfile.ZipFile(source) as z:
 members={n:z.read(n) for n in z.namelist()}
 styles=ET.fromstring(members['xl/styles.xml']);xfs=styles.find('m:cellXfs',ns);wrapped={}
 def wrap_style(idx):
  if idx not in wrapped:
   xf=copy.deepcopy(xfs[idx]);al=xf.find('m:alignment',ns)
   if al is None:al=ET.SubElement(xf,N+'alignment')
   al.set('wrapText','1');xf.set('applyAlignment','1');wrapped[idx]=len(xfs);xfs.append(xf)
  return wrapped[idx]
 for i,s in enumerate(src,1):
  key=f'xl/worksheets/sheet{i}.xml';tree=ET.fromstring(members[key]);rows=tree.find('m:sheetData',ns)
  cells={c.attrib['r']:c for r in rows for c in r}
  for addr,value in inputs.get(s.title,{}).items():
   assert s[addr].fill.fgColor.rgb=='FFFFF2CC',(s.title,addr,'not yellow')
   c=cells[addr]
   for child in list(c):
    if child.tag in [N+'v',N+'is',N+'f']:c.remove(child)
   c.attrib.pop('t',None)
   val=ex[s.title][addr].value
   assert val==value,(s.title,addr,val,value)
   if val is not None:
    if isinstance(val,str):
     c.set('t','inlineStr');is_=ET.SubElement(c,N+'is');ET.SubElement(is_,N+'t').text=val
     c.set('s',str(wrap_style(int(c.get('s','0')))))
    else:ET.SubElement(c,N+'v').text=str(val)
  for row in s:
   for cell in row:
    if cell.data_type=='f':
     c=cells[cell.coordinate];v=c.find('m:v',ns)
     if v is None:v=ET.SubElement(c,N+'v')
     value=ex[s.title][cell.coordinate].value
     if isinstance(value,str):c.set('t','str');v.text=value
     elif value is None:v.text=None;c.attrib.pop('t',None)
     else:c.attrib.pop('t',None);v.text=str(value)
  for row in rows:
   if row.attrib['r'] in layout.get(s.title,{}):row.set('ht',str(layout[s.title][row.attrib['r']]));row.set('customHeight','1')
  members[key]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
 xfs.set('count',str(len(xfs)));members['xl/styles.xml']=ET.tostring(styles,encoding='utf-8',xml_declaration=True)
 with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as out:
  for n,b in members.items():out.writestr(n,b)

final=openpyxl.load_workbook(dest);cached=openpyxl.load_workbook(dest,data_only=True)
checks=[]
def ok(name,condition):
 assert condition,name
 checks.append(name)
ok('7 sheets and order unchanged',src.sheetnames==final.sheetnames)
changed=[];formula_count=0
for a,b in zip(src,final):
 ok(a.title+' validations preserved',str(a.data_validations)==str(b.data_validations))
 with zipfile.ZipFile(source) as za,zipfile.ZipFile(dest) as zb:
  sheetpath=f'xl/worksheets/sheet{src.sheetnames.index(a.title)+1}.xml'
  ta,tb=ET.fromstring(za.read(sheetpath)),ET.fromstring(zb.read(sheetpath))
  ok(a.title+' conditional formatting preserved',[ET.tostring(e) for e in ta.findall('m:conditionalFormatting',ns)]==[ET.tostring(e) for e in tb.findall('m:conditionalFormatting',ns)])
 ok(a.title+' merges/freeze preserved',str(a.merged_cells)==str(b.merged_cells) and a.freeze_panes==b.freeze_panes)
 for row in a:
  for c in row:
   target=b[c.coordinate]
   if c.data_type=='f':
    formula_count+=1;ok(a.title+'!'+c.coordinate+' formula unchanged',target.value==c.value)
    ok(a.title+'!'+c.coordinate+' cache present',cached[a.title][c.coordinate].value is not None)
   if c.value!=target.value:
    ok(a.title+'!'+c.coordinate+' only yellow input changed',c.fill.fgColor.rgb=='FFFFF2CC' and c.coordinate in inputs.get(a.title,{}))
    changed.append(a.title+'!'+c.coordinate)
   if not(c.coordinate in inputs.get(a.title,{}) and isinstance(inputs[a.title][c.coordinate],str)):
    ok(a.title+'!'+c.coordinate+' style preserved',c._style==target._style)
cost=cached['1_Cost_Job'];price=cached['2_Pricing'];ch=cached['4_Channel_Fit']
expect={'1_Cost_Job!B27':.02225,'1_Cost_Job!B28':.03,'1_Cost_Job!B54':22.5,'1_Cost_Job!B65':104.53,'1_Cost_Job!B66':.0681625,'1_Cost_Job!B67':.1306625,'2_Pricing!B7':.2044875,'2_Pricing!B21':.82959375,'2_Pricing!B33':.3408125,'4_Channel_Fit!B9':1592.82,'4_Channel_Fit!B16':.15625,'4_Channel_Fit!B22':5000}
for a,v in expect.items():
 sn,addr=a.split('!');ok(a+' arithmetic',math.isclose(cached[sn][addr].value,v,rel_tol=1e-10,abs_tol=1e-10))
for row,r in zip(range(39,45),[.5,.6,.7,.8,.9,.95]):
 ok('pricing sensitivity row '+str(row),math.isclose(price[f'B{row}'].value,.05453/r,rel_tol=1e-9) and math.isclose(price[f'C{row}'].value,1-.05453/r/.4,rel_tol=1e-9))
for addr,v in {'B34':20,'C34':13,'D34':10}.items():ok('channel score '+addr,ch[addr].value==v)
stress=json.loads((Q/'engine-stress.json').read_text());q=.05*3/60*9;api=(5000*1.25+2*5000*.1+3*1000+3*800*5)/1e6;per=api*1.08+.008+q
for case in stress['stress']:
 if case['case']=='containment':
  r=case['rate'];ok('engine completion '+str(r),math.isclose(case['cost'][0][0],per/r) and math.isclose(case['cost'][1][0],(per+.05)/r))
 elif case['case']=='variantB':ok('engine variant B',math.isclose(case['cost'][0][0],(per+.2*6/60*9)/.8))
 elif case['case']=='batch':ok('engine batch',math.isclose(case['cost'][0][0],(api*.5*1.08+.008+q)/.8))
 elif case['case']=='overhead2x':ok('engine overhead2x',math.isclose(case['cost'][1][0],(per+.1)/.8))
 else:ok('zero containment existing guard observed',case['cost']==[[0],[0]])
extra={'no_cache':{'direct_cost':(.03*1.08+.008+q)/.8,'full_cost':(.03*1.08+.008+q+.05)/.8},'half_volume_same_rate':{'attempt':500,'completed':400,'ARPU_each':80,'full_cost':(500*per+50)/400,'full_GM':1-(500*per+50)/160},'payment_fee3pct_sensitivity':{'full_GM':1-cost['B67'].value/.4-.03},'minimum_r_direct_60pct':per/.16,'minimum_r_full_60pct':(per+.05)/.16,'minimum_r_full_50pct':(per+.05)/.2,'overhead_fullGM50_at80pct':1000*(.8*.2-per)}
report={'status':'PASS','checks_count':len(checks),'formulas_preserved':formula_count,'changed_inputs':changed,'packaging':'Original OOXML preserved; artifact-authored inputs/caches transferred; only yellow text wrapping and local row heights changed.','export_findings':repair,'engine':'artifact-tool; Microsoft Excel desktop not executed','extra_independent_scenarios':extra,'limitations':['zero completed cost undefined although template IF returns0','no live AI eval or human stranger test','financial drivers are planning estimates']}
(Q/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')

# Compact the template fields into one A4 page, with cell-level traceability.
for name,file in [('VN','arial.ttf'),('VNB','arialbd.ttf')]:pdfmetrics.registerFont(TTFont(name,'C:/Windows/Fonts/'+file))
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyVN',fontName='VN',fontSize=9.2,leading=12,spaceAfter=4));styles.add(ParagraphStyle(name='SmallVN',fontName='VN',fontSize=8,leading=10));styles.add(ParagraphStyle(name='HeadVN',fontName='VNB',fontSize=10.5,leading=14,spaceBefore=5,spaceAfter=4));styles.add(ParagraphStyle(name='TitleVN',fontName='VNB',fontSize=17,leading=20,spaceAfter=5))
story=[]
def para(t,style='BodyVN'):return Paragraph(t,styles[style])
def add(t,style='BodyVN'):story.append(para(t,style))
def table(rows,widths):
 t=Table([[para(str(c),'SmallVN') for c in row] for row in rows],colWidths=widths,hAlign='LEFT');t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EAF6')),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#667085')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]));story.append(t);story.append(Spacer(1,4))
add('AI Notes - Monetization One-Pager','TitleVN')
add('Dương Thị Hồng Viên | MSSV 2A202602385 | Ngày 6 / Day22 | 08/10/2026','SmallVN')
add('Usage theo nháp hoàn thành - PLG. Kịch bản kinh doanh cho prototype Option B; dữ liệu hiện dựng sẵn, chưa có model live hoặc khách trả phí. USD; chi phí vận hành bên dưới là ước tính.','SmallVN')
add('1. Pricing','HeadVN')
add('<b>Buyer/ngân sách:</b> Trưởng đào tạo/chủ SME duyệt, người học sử dụng (giả thuyết). A: công cụ phần mềm hỗ trợ ghi chú. B: giảm công chuẩn bị nháp từ ngân sách vận hành/nhân sự - tôi chọn B để thử, không thay toàn bộ giảng viên. <b>Job:</b> nháp cá nhân/bài đủ ý chính, chỗ chưa hiểu, thực hành và nguồn hợp lệ; lỗi/trùng/retry không tính phí [Cost B5-B6].')
add('<b>Metric:</b> Attribution 1/10, Autonomy 0/10 [Metric B10/B18] nên model gợi ý Seat/Hybrid [B21]. Tôi thử Usage vì tần suất học khác nhau, giữ quyền sửa miễn phí; cần metering/eval trước bán [B30:B34]. Cap 400 job/tháng/khách [Pricing B11], giá $0,40/job [B19], tạo lại phải đồng ý lượt mới, tối đa 3 attempt/job trong pilot (kiểm soát đề xuất).')
table([['Chỉ số','Kết quả và ô Excel'],['Cost/job và giá sàn trực tiếp','$0,06816 [Pricing B5]; 3x = $0,20449 [B7]'],['Giá / GM trực tiếp','$0,40 [Pricing B19] / 82,96% [B21]'],['Cost và biên sau overhead','$0,13066 [Cost B67]; GM = 1 - B67/Pricing B19 = 67,33%'],['Containment giả định / ngưỡng GM60%','80% [Cost B10], chưa eval. Direct 34,08% [Pricing B33]; gồm overhead 65,33% = (Cost B62+B63+B64)/B9 / (Pricing B19*(1-B32))']], [218,321])
add('<b>Neo giá:</b> tiết kiệm giả định $800/tháng [Pricing B10] từ 20phút/job x $6/giờ x 400; thu20% giá trị. Trần25% = $0,50/job [B13]; salary $1000 [B14] chỉ đối chiếu, không phải nhân sự bị thay. <b>Gãy:</b> full GM&lt;50% nếu completion&lt;52,265% = Cost B65/B9/(Pricing B19*0,5). Cache phải đúng prefix/TTL; overhead tăng đôi làm full GM51,71%.')
add('<b>Benchmark</b> [Metric A26:D27]: Notion Agents $10/1000credits khi trả tháng (notion.com/help/what-are-notion-credits); Otter Pro $16,99/user/tháng + quota (otter.ai/pricing). Kiểm tra08/10/2026; credit/phút khác job ghi chú.','SmallVN')
add('2. Go-to-market','HeadVN')
add('<b>PLG</b> [Channel B38]: tự thử nháp có nguồn rồi mua theo lượt. ARPU $160 [B5] chỉ khi dùng đủ400job, ACV $1920 [B10]; CAC budget $1592,82 [B9] = ARPU x directGM x 12tháng [B8]. AE quota giả định $75.000/năm [B13] cho0,15625deal/ngày [B16]: không bị loại bằng công suất. CPO $1000 [B20]/win20% [B21] = CAC $5000 [B22], cao3,14x ngân sách [B23]. Các driver sales là ước tính, không phải benchmark đã đo.')
add('<b>Pain:</b> 21:30, vừa học xong cần gom ghi chú làm lab hôm sau trong VLearn [Plan B5:B7]. <b>Điểm nhúng:</b> sidebar bài VLearn mới là đề xuất, chưa có quyền/API [B8]. Trước mắt nhập transcript có phép vào prototype; phải đo ma sát này. CAC PLG dự kiến $200/5paid=$40, chưa đo; giảm volume một nửa làm full GM51,71%.')
table([['Kế hoạch - Viên phụ trách','Việc và KPI (Plan B12:D18)'],['Tháng 1 - học','2 SME pilot; phỏng vấn buyer, eval50ca, đo cost/thời gian, xác minh tích hợp. Chưa tuyển.'],['Tháng 2-3 - đòn bẩy','Mục tiêu5paid; thử onboarding/metering/cap, đo CAC/retention; gate completion>=80% và fullGM>=60%.'],['Tháng 4+ - có điều kiện','Mục tiêu10 SME ở1ngách gần, chỉ sau khi qua gate và có retention.']], [160,379])
add('3. Evidence pack','HeadVN')
add('<b>Eval:</b> chưa có live, Viên đo50ca trước22/10/2026 [Plan B23:D23]. <b>Procurement Q&amp;A:</b> có dự thảo sai nguồn/data training/xuất-xóa; xác minh trước15/10/2026 [B24:D24]. <b>Pilot:</b> dự kiến2SME,2tuần, báo cáo06/11/2026 [B25:D25]; chưa tuyển/chạy. Không dùng3tester usability cũ làm containment.')
add('<b>Test người lạ:</b> chưa thực hiện [Plan B29:B32]; Viên dự kiến10/10/2026. Cho nhóm khác đọc2phút, trả lời bán gì/cho ai/đơn vị; có lãi vì số nào; tiếp cận qua đâu/vì sao. Số câu hỏi lại còn trống, mục tiêu&lt;=3. Chi tiết nguồn, phạm vi chi phí và test tại assumptions-and-sources.md, qa/verification.json.','SmallVN')
pdf=O/'DuongThiHongVien_Day22_onepager.pdf';SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=28,leftMargin=28,topMargin=25,bottomMargin=25).build(story)
reader=PdfReader(pdf);ok('PDF exactly 1 page',len(reader.pages)==1);ok('PDF identity extracted','2A202602385' in reader.pages[0].extract_text())
doc=pypdfium2.PdfDocument(str(pdf));doc[0].render(scale=2).to_pil().save(Q/'onepager.png')
report['pdf_pages']=len(reader.pages);report['checks_count']=len(checks);(Q/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'checks':len(checks),'formula_count':formula_count,'pdf_pages':len(reader.pages),'expected':expect,'export_findings':repair},ensure_ascii=False))
