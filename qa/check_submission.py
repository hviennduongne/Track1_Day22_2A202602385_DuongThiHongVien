from pathlib import Path
import json, math, hashlib
import openpyxl
from pypdf import PdfReader

p = Path(__file__).resolve().parents[1]
inputs = json.loads((p / 'inputs.json').read_text(encoding='utf-8'))
file = p / 'output/DuongThiHongVien_Day22_model.xlsx'
w = openpyxl.load_workbook(file)
v = openpyxl.load_workbook(file, data_only=True)
original = openpyxl.load_workbook(p / 'templates/model.xlsx')
checks = []
def check(label, ok):
    checks.append({'check': label, 'pass': bool(ok)})
    assert ok, label
check('Original sheet order and count', w.sheetnames == original.sheetnames)
for name, cells in inputs.items():
    for address, expected in cells.items():
        check(f'Input {name}!{address}', w[name][address].value == expected)
for s in original:
    for row in s:
        for c in row:
            if c.data_type == 'f':
                check(f'Formula {s.title}!{c.coordinate}', w[s.title][c.coordinate].value == c.value)
                check(f'Cached result {s.title}!{c.coordinate}', v[s.title][c.coordinate].value is not None)
cost = v['1_Cost_Job']
check('Completed denominator', cost['B11'].value == 800)
check('Five-part total', math.isclose(cost['B65'].value, 104.53))
check('Cost per completed including overhead', math.isclose(cost['B67'].value,104.53/800))
check('CAC budget', math.isclose(v['4_Channel_Fit']['B9'].value,1592.82))
check('One channel', v['4_Channel_Fit']['B38'].value == 'PLG')
pdf = p / 'output/DuongThiHongVien_Day22_onepager.pdf'
r = PdfReader(pdf)
check('One-page PDF',len(r.pages)==1)
t = r.pages[0].extract_text()
for phrase in ['2A202602385','6300','31500','1592','chưa thực hiện']:
    check('PDF contains '+phrase,phrase in t)
for name in ['ai-log.md','reflection.md','critique-log.md','requirements-audit.md','submission-status.md','evidence/channel-evidence.md','evidence/eval-plan.md','evidence/eval-results.md','evidence/procurement-qa.md','evidence/pilot-plan.md','evidence/number-traceability.md','evidence/stranger-test.md']:
    check('Submission asset '+name,(p/name).is_file() and (p/name).stat().st_size>100)
check('Human test not invented',v['5_90Day_Plan']['B32'].value is None)
report = {'status':'PASS_FILE_CHECKS','checks':len(checks),'results':checks,'sha256':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in [file,pdf]},'remaining':'Human stranger test has no actual results. Containment and commercial drivers are estimates; quality eval/pilot remain planned.'}
(p/'qa/final-submission-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'status':report['status'],'checks':len(checks),'sha256':report['sha256']},ensure_ascii=False))
